import re, html
from docx import Document
from docx.shared import Pt

src = open('interview-prep.html', encoding='utf-8').read()
# split into sections by <section id=...><h2>title</h2>
secs = re.findall(r'<section id="([^"]+)".*?<h2>(.*?)</h2>(.*?)</section>', src, re.S)
doc = Document()
style = doc.styles['Normal']
style.font.size = Pt(11)
doc.add_heading('1 YOE Full-Stack Interview Prep (MERN / .NET) - Full Q&A', 0)
for sid, title, body in secs:
    doc.add_heading(html.unescape(re.sub(r'<.*?>', '', title)).strip(), 1)
    items = re.findall(r'<details.*?><summary>(.*?)</summary>(.*?)</details>', body, re.S)
    if not items:
        txt = html.unescape(re.sub(r'<.*?>', ' ', body)).strip()
        if txt:
            doc.add_paragraph(txt)
    for s, b in items:
        q = html.unescape(re.sub(r'<.*?>', '', s)).strip()
        doc.add_heading(q, 2)
        # code blocks
        for m in re.finditer(r'<pre>(.*?)</pre>', b, re.S):
            code = html.unescape(re.sub(r'<.*?>', '', m.group(1))).strip()
            doc.add_paragraph(code, style='No Spacing').runs[0].font.name = 'Consolas'
        rest = re.sub(r'<pre>.*?</pre>', ' ', b, flags=re.S)
        t = html.unescape(re.sub(r'\s+', ' ', re.sub(r'<.*?>', ' ', rest))).strip()
        if t:
            doc.add_paragraph(t)
doc.save('Interview-Prep-Full.docx')
print('saved Interview-Prep-Full.docx')
