import os, re
from docx import Document
from docx.shared import Pt, RGBColor
BASE = r"C:\Users\Qsp\Documents\mernstack_interview_questions\mern-200"
OUT = os.path.join(BASE, "word")
os.makedirs(OUT, exist_ok=True)
SUBJECTS = ["js-frontend","react-redux","express-node","mongodb","dsa-git"]
pat = re.compile(r"^\s*(\d+)\.\s+\*\*(.+?)\*\*\s*(.*)$")
def add_answer(doc, text):
    # split `code` into runs
    parts = re.split(r"(`[^`]+`)", text)
    p = doc.add_paragraph()
    p.style.font.size = Pt(11)
    for part in parts:
        if not part: continue
        if part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1])
            r.font.name = "Consolas"; r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(0x1F,0x6F,0xEB)
            r.font.bold = True
        else:
            r = p.add_run(part)
            r.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(6)
for slug in SUBJECTS:
    md = os.path.join(BASE, slug, "README.md")
    lines = open(md, encoding="utf-8").read().splitlines()
    title = lines[0].lstrip("# ").strip()
    items=[]; q=None; buf=[]
    for line in lines[1:]:
        s=line.strip()
        if s.startswith("##") or s.startswith("> "): continue
        m=pat.match(line)
        if m:
            if q: items.append((q, " ".join(buf).strip()))
            q=m.group(2).strip(); rest=m.group(3).strip(); buf=[rest] if rest else []
        elif q and s:
            buf.append(s)
    if q: items.append((q, " ".join(buf).strip()))
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(title + " — Detailed (200)", 0)
    for i,(qq,aa) in enumerate(items,1):
        doc.add_heading(f"{i}. {qq}", 2)
        add_answer(doc, aa)
    out = os.path.join(OUT, f"{slug}-200-detailed.docx")
    doc.save(out)
    print(f"{slug}: {len(items)} -> {out}")
# combined
doc = Document()
doc.styles["Normal"].font.size = Pt(11)
doc.add_heading("MERN Fullstack — 1000 Q&A Detailed", 0)
for slug in SUBJECTS:
    md = os.path.join(BASE, slug, "README.md")
    lines = open(md, encoding="utf-8").read().splitlines()
    title = lines[0].lstrip("# ").strip()
    doc.add_page_break()
    doc.add_heading(title, 1)
    items=[]; q=None; buf=[]
    for line in lines[1:]:
        s=line.strip()
        if s.startswith("##") or s.startswith("> "): continue
        m=pat.match(line)
        if m:
            if q: items.append((q, " ".join(buf).strip()))
            q=m.group(2).strip(); rest=m.group(3).strip(); buf=[rest] if rest else []
        elif q and s:
            buf.append(s)
    if q: items.append((q, " ".join(buf).strip()))
    for i,(qq,aa) in enumerate(items,1):
        doc.add_heading(f"{i}. {qq}", 2)
        add_answer(doc, aa)
combined = os.path.join(OUT, "mern-1000-detailed-full.docx")
doc.save(combined)
print(f"combined -> {combined}")
